"""Build the Week 6 DOCX report from the executed notebook outputs."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "outputs"
RESULTS_PATH = OUTPUT_DIR / "week6_results.json"
REPORT_PATH = ROOT / "Week6_Customer_Support_Intent_Classification_Report.docx"


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_code_block(document, code):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(35, 55, 75)


def format_metric(value):
    return "N/A" if value is None else f"{value:.4f}"


def main():
    if not RESULTS_PATH.exists():
        raise FileNotFoundError(
            f"{RESULTS_PATH} is missing. Execute all cells in the Week 6 notebook first."
        )
    results_data = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    metrics = results_data["metrics"]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.color.rgb = RGBColor(26, 82, 91)
    for heading_name in ("Heading 1", "Heading 2"):
        styles[heading_name].font.name = "Aptos Display"
        styles[heading_name].font.color.rgb = RGBColor(26, 82, 91)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Week 6 Assignment\nCustomer Support Intent Classification")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "AI-Based Smart Queue & Time Management System | TECH 405\n"
        "Bishwadeep Rai | Westcliff University, California | Professor Aashish Dhakal\n"
        "Feed-forward Neural Network (MLP) vs. PyTorch LSTM"
    )

    document.add_heading("Executive summary", level=1)
    document.add_paragraph(
        "This classification experiment predicts the service intent of an incoming "
        "customer-support request, a practical request-routing component for the "
        "Smart Queue project. The executed notebook compares two distinct neural "
        "models on identical held-out data: a TF-IDF feed-forward MLP and a "
        "token-embedding recurrent LSTM. The target is the dataset's labeled intent."
    )
    best_model = max(
        ("Feed-forward MLP", "PyTorch LSTM"),
        key=lambda name: metrics[name]["f1_macro"],
    )
    document.add_paragraph(
        f"By held-out macro F1, the stronger model in this run was {best_model}. "
        "Read the full metric table and figures below; macro scores weight each "
        "intent equally."
    )

    document.add_heading("Dataset and ground truth", level=1)
    source = document.add_paragraph()
    source.add_run("Dataset: ")
    add_hyperlink(
        source,
        "Bitext Customer Service Tagged Training Dataset",
        results_data["dataset_url"],
    )
    document.add_paragraph(
        f"Source-reported size is approximately 26,872 request/response examples. "
        f"The loaded CSV contains {results_data['source_rows']:,} rows and "
        f"{results_data['number_of_intents']} intent classes. The input is the "
        "customer's `instruction`; `intent` is the ground-truth class used for "
        "training and validation."
    )
    document.add_paragraph(
        "The source metadata declares CDLA-Sharing-1.0. Its dataset card identifies "
        "the corpus as hybrid synthetic data produced with NLP/NLG and automated "
        "labeling. These labels provide a reference target for this assignment, "
        "but they are not independently verified real-world queue outcomes. "
        "Review the linked license terms before redistributing source data."
    )
    document.add_paragraph(
        "Leakage controls: the response and correlated `category` field are excluded; "
        "the vocabulary/vectorizer is fitted only on training instructions; and "
        "train/validation/test records are mutually exclusive and stratified."
    )

    dataset_table = document.add_table(rows=1, cols=2)
    dataset_table.style = "Light Shading Accent 1"
    dataset_table.rows[0].cells[0].text = "Partition"
    dataset_table.rows[0].cells[1].text = "Examples"
    for label, key in (
        ("Training", "train_rows"),
        ("Validation", "validation_rows"),
        ("Held-out test", "test_rows"),
    ):
        cells = dataset_table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{results_data[key]:,}"

    document.add_heading("Models and evaluation method", level=1)
    document.add_paragraph(
        "Model 1 — Feed-forward MLP: training-only TF-IDF unigram/bigram features "
        "are passed through two ReLU dense layers (256 and 128 units), dropout, "
        "and a 27-class output. Sparse batches are densified only while training."
    )
    document.add_paragraph(
        "Model 2 — PyTorch LSTM: a training-only word vocabulary maps the request "
        "to token IDs; a 64-dimensional embedding feeds a 64-unit recurrent LSTM, "
        "dropout, and a 27-class output. Packed sequences prevent padding from "
        "being treated as request content."
    )
    document.add_paragraph(
        f"Split: {results_data['train_rows']:,} train / "
        f"{results_data['validation_rows']:,} validation / "
        f"{results_data['test_rows']:,} test, stratified with random seed "
        f"{results_data['seed']}. Validation macro F1 selects the checkpoint. "
        "The majority baseline predicts the most common training intent."
    )
    document.add_paragraph(
        "Precision, recall, and F1 are macro averages across intent classes. "
        "ROC-AUC is macro one-vs-rest multiclass AUC. The test set is used once "
        "for final reporting, not for model selection."
    )

    document.add_heading("Held-out test results", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    headers = [
        "Model",
        "Accuracy",
        "Precision\n(macro)",
        "Recall\n(macro)",
        "F1\n(macro)",
        "ROC-AUC\n(OvR macro)",
    ]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for name, metric in metrics.items():
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = format_metric(metric["accuracy"])
        cells[2].text = format_metric(metric["precision_macro"])
        cells[3].text = format_metric(metric["recall_macro"])
        cells[4].text = format_metric(metric["f1_macro"])
        cells[5].text = format_metric(metric["roc_auc_ovr_macro"])

    document.add_heading("Result snapshots", level=1)
    for image_name, caption in (
        ("confusion_matrices.png", "Figure 1. Row-normalized held-out confusion matrices."),
        ("roc_curves.png", "Figure 2. Micro-averaged one-vs-rest ROC curves."),
        ("training_loss.png", "Figure 3. Training cross-entropy by epoch."),
    ):
        image_path = OUTPUT_DIR / image_name
        if not image_path.exists():
            raise FileNotFoundError(
                f"{image_path} is missing. Re-run all notebook cells before building the report."
            )
        document.add_picture(str(image_path), width=Inches(6.3))
        caption_paragraph = document.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Critical code segments", level=1)
    document.add_paragraph(
        "1. The intent label is split before any trainable text transformation. "
        "This isolates final testing and fits vocabulary statistics on training text only:"
    )
    add_code_block(
        document,
        "train_idx, holdout_idx = train_test_split(\n"
        "    all_indices, test_size=0.30, random_state=SEED, stratify=y\n"
        ")\n"
        "valid_idx, test_idx = train_test_split(\n"
        "    holdout_idx, test_size=0.50, random_state=SEED,\n"
        "    stratify=y[holdout_idx]\n"
        ")\n"
        "X_train_tfidf = vectorizer.fit_transform(texts[train_idx])\n"
        "X_test_tfidf = vectorizer.transform(texts[test_idx])"
    )
    document.add_paragraph(
        "2. The LSTM processes only each request's real token sequence and maps "
        "the final hidden state to intent scores:"
    )
    add_code_block(
        document,
        "embedded = self.embedding(token_ids)\n"
        "packed = pack_padded_sequence(\n"
        "    embedded, lengths.cpu(), batch_first=True, enforce_sorted=False\n"
        ")\n"
        "_, (hidden, _) = self.lstm(packed)\n"
        "return self.output(self.dropout(hidden[-1]))"
    )

    document.add_heading("Comparison and conclusion", level=1)
    mlp_scores = metrics["Feed-forward MLP"]
    lstm_scores = metrics["PyTorch LSTM"]
    delta = lstm_scores["f1_macro"] - mlp_scores["f1_macro"]
    if abs(delta) < 1e-12:
        conclusion = "The two models achieved the same held-out macro F1 in this run."
    else:
        winner = "LSTM" if delta > 0 else "MLP"
        conclusion = (
            f"The {winner} achieved the higher held-out macro F1 "
            f"({max(mlp_scores['f1_macro'], lstm_scores['f1_macro']):.4f}); "
            f"the difference between models was {abs(delta):.4f}."
        )
    document.add_paragraph(
        conclusion
        + " Consider precision/recall trade-offs and the per-intent confusion "
        "matrices before selecting a routing model. The baseline helps show whether "
        "either neural model is useful beyond predicting the most common class."
    )
    document.add_paragraph(
        "This result generalizes only to this fixed synthetic dataset split. "
        "Template reuse can make random held-out examples easier than truly new "
        "customer phrasing. A real queue deployment requires external validation "
        "on human-reviewed, representative requests, monitoring, privacy safeguards, "
        "and a safe human-escalation path."
    )

    document.add_heading("Reproduction and publication", level=1)
    document.add_paragraph(
        "From the repository root, install the assignment dependencies and launch "
        "the notebook:"
    )
    add_code_block(
        document,
        "python -m pip install -r week6assignment/requirements.txt\n"
        "jupyter notebook week6assignment/week6_customer_support_intent_classification.ipynb\n"
        "python week6assignment/build_report.py"
    )
    github_paragraph = document.add_paragraph()
    github_paragraph.add_run("Course repository: ")
    add_hyperlink(
        github_paragraph,
        "BishwadeepRai/Smart-queue-Time-Management-system",
        "https://github.com/BishwadeepRai/Smart-queue-Time-Management-system",
    )
    notebook_paragraph = document.add_paragraph()
    notebook_paragraph.add_run("Week 6 notebook permalink: ")
    add_hyperlink(
        notebook_paragraph,
        "week6_customer_support_intent_classification.ipynb",
        "https://github.com/BishwadeepRai/Smart-queue-Time-Management-system/blob/main/week6assignment/week6_customer_support_intent_classification.ipynb",
    )
    document.add_paragraph(
        "The assignment files must be pushed before the notebook permalink can be "
        "submitted as a public GitHub link. Do not add the downloaded data cache "
        "to the repository."
    )

    document.save(REPORT_PATH)
    print(f"Created report: {REPORT_PATH}")


if __name__ == "__main__":
    main()
