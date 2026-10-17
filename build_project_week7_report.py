"""Build the project-wide Week 7 technical report from saved experiment results."""

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT_PATH = ROOT / "Project_Week7_Technical_Report.docx"
GITHUB_URL = "https://github.com/BishwadeepRai/Smart-queue-Time-Management-system"
CALL_DATA_URL = "https://www.kaggle.com/datasets/jpmiller/call-centre-queue-simulation"
BITEXT_DATA_URL = "https://huggingface.co/datasets/bitext/Bitext-customer-support-llm-chatbot-training-dataset"


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def add_figure(document, path, caption):
    if not path.is_file():
        raise FileNotFoundError(f"Required report figure is missing: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.1))
    add_caption(document, caption)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_paragraph()
    return table


def add_code(document, code):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def fmt(value):
    return "Not available" if value is None else f"{value:.4f}"


def fmt_auc(value):
    return "Not available" if value is None else f"{value:.5f}"


def first_image(pattern):
    matches = sorted((ROOT / "week 2" / "images").glob(pattern))
    if not matches:
        raise FileNotFoundError(f"No Week 2 image matches {pattern!r}")
    return matches[0]


def main():
    w4 = json.loads((ROOT / "week 4" / "results.json").read_text(encoding="utf-8"))
    w5 = json.loads((ROOT / "week 5" / "outputs" / "week5_results.json").read_text(encoding="utf-8"))
    w6 = json.loads((ROOT / "week 6" / "outputs" / "week6_results.json").read_text(encoding="utf-8"))

    w4m = w4["test_metrics"]
    w5m = w5["test_metrics"]
    w6m = w6["metrics"]
    w4_selected = next(
        item for item in w4["experiments"] if item["id"] == w4["best_experiment_id"]
    )
    w5_selected = w5["selected_parameters"]
    w5macro = w5m["macro_avg"]
    w6_mlp = w6m["Feed-forward MLP"]
    w6_lstm = w6m["PyTorch LSTM"]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.line_spacing = 1.5
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Times New Roman"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(100)
    title_run = title.add_run(
        "Smart Queue & Time Management System\n"
        "Project-Wide Week 7 Technical Report"
    )
    title_run.bold = True
    title_run.font.size = Pt(18)
    for line in (
        "Bishwadeep Rai",
        "Westcliff University",
        "TECH 405: Artificial Neural Network and Deep Learning",
        date.today().strftime("%B %d, %Y"),
    ):
        paragraph = document.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "This report documents the project's progression across Weeks 2, 4, 5, "
        "and 6. Weeks 2, 4, and 5 use a simulated call-centre queue dataset, "
        "moving from exploratory analysis of waiting time to two separate "
        "binary service-standard classification experiments. Week 6 is a "
        "distinct multiclass customer-support intent-routing experiment using "
        "the Bitext dataset. The report preserves these different prediction "
        "problems rather than presenting them as one model or one continuous "
        "target. It summarizes the datasets, exploratory analysis, preprocessing, "
        "model development, held-out evaluation, figures, and methodological "
        "limitations based on the project artifacts. The call-centre data is "
        "simulated, and the Week 6 data is synthetic/automatically labeled; "
        "none of the results establish operational readiness."
    )

    document.add_heading("Project Scope and Problem Statements", level=1)
    document.add_paragraph(
        "The coursework evolves over several weekly submissions, but its later "
        "classification tasks are not interchangeable. Week 2 states the intended "
        "system objective as estimating queue waiting time (`wait_length`) to "
        "support wait-now-or-return decisions. Week 4 instead predicts whether "
        "a call meets the 60-second service standard, and Week 5 predicts the "
        "next call's same binary outcome from a recent sequence. Week 6 changes "
        "to routing customer text among 27 support intents. These are separately "
        "evaluated tasks with their own targets, split strategies, models, and "
        "results."
    )
    add_table(
        document,
        ["Stage", "Dataset and target", "Task / model"],
        [
            ["Week 2", "Simulated call-centre CSV; `wait_length` is the stated future target.",
             "Exploratory analysis; no trained model in this notebook."],
            ["Week 4", "Same call-centre CSV; `meets_standard`.",
             "Binary feed-forward neural-network classification."],
            ["Week 5", "Same call-centre CSV; next-call `meets_standard`.",
             "Binary sequence classification with a PyTorch LSTM."],
            ["Week 6", "Bitext support requests; `intent`.",
             "27-class text classification with TF-IDF MLP and PyTorch LSTM."],
        ],
    )

    document.add_heading("Datasets and Provenance", level=1)
    document.add_paragraph(
        "Weeks 2, 4, and 5 use the repository's 51,708-row, nine-column "
        "`simulated_call_centre.csv` dataset. The Week 2 dataset card records "
        "2021-01-01 through 2021-12-31, 261 Monday-to-Friday operating days, "
        "and identifies the data as generated by a discrete-event simulation. "
        f"Source: {CALL_DATA_URL}. The CSV is shared within this repository."
    )
    document.add_paragraph(
        "Week 6 uses the Bitext Customer Service Tagged Training Dataset, with "
        f"{w6['source_rows']:,} labeled requests across {w6['number_of_intents']} "
        f"intents. Its documented source is {BITEXT_DATA_URL}. The raw file is "
        "downloaded to a local cache by the Week 6 workflow rather than stored "
        "in the repository."
    )
    document.add_paragraph(
        f"Repository: {GITHUB_URL}. Dataset source pages and repository links "
        "are recorded for reproducibility; redistribution remains subject to "
        "the respective dataset terms."
    )

    document.add_heading("Exploratory Data Analysis", level=1)
    document.add_paragraph(
        "Week 2 audited data types, missingness, distributions, the target, "
        "and temporal patterns. The majority of records have zero waiting time, "
        "so `wait_length` is zero-inflated and right-skewed. The dataset is "
        "chronologically ordered and temporal patterns motivated respecting "
        "time order in the later call-centre experiments. For the classification "
        "target used in Weeks 4 and 5, 4,227 of 51,708 calls (8.17%) missed the "
        "standard, making class imbalance an important evaluation consideration."
    )
    add_figure(
        document,
        first_image("figure-02*"),
        "Week 2: call-centre target distribution as saved in the executed EDA notebook.",
    )
    add_figure(
        document,
        first_image("figure-05*"),
        "Week 2: temporal EDA figure from the executed notebook.",
    )
    document.add_paragraph(
        "Week 4's recorded EDA includes class distribution, queue position by "
        "outcome, temporal service-standard rates, and feature correlations. "
        "Week 6's audit found no missing or blank values in its five source "
        "fields; intent counts range from 950 to 1,000 examples. Case- and "
        "whitespace-normalization identified 1,187 repeated-request groups "
        "(2,598 additional rows), and all normalized duplicates had consistent "
        "intent labels."
    )
    add_figure(
        document,
        ROOT / "week 4" / "images" / "fig1_class_distribution.png",
        "Week 4: imbalance in the binary service-standard outcome.",
    )
    add_figure(
        document,
        ROOT / "week 6" / "images" / "intent_class_distribution_and_duplicates.png",
        "Week 6: intent-class counts and normalized duplicate-request analysis.",
    )

    document.add_heading("Model Development and Methodology", level=1)
    document.add_heading("Week 4: Feed-Forward Neural Network", level=2)
    document.add_paragraph(
        "The binary target is `meets_standard` (False = missed, True = met). "
        "The runner sorts records by call start time and makes a chronological "
        "70/15/15 train/validation/test split. Eight arrival/calendar features "
        "are used. `call_id`, `wait_length`, `call_answered`, `call_ended`, and "
        "`service_length` are excluded; in particular, `wait_length <= 60` "
        "determines the label and would directly leak the target. Four Keras "
        "feed-forward configurations are compared. The selected configuration "
        "is chosen on validation missed-class F1, then validation macro F1; "
        "the test set is used only for final metrics."
    )
    document.add_paragraph(
        f"The selected run used hidden layers {w4_selected['architecture']}, "
        f"learning rate {w4_selected['learning_rate']}, batch size "
        f"{w4_selected['batch_size']}, and dropout {w4_selected['dropout']}. "
        "Training-set class weights address imbalance; the chronological test "
        "set is kept separate from model selection."
    )
    add_code(
        document,
        'best = max(experiment_results, key=lambda r: '
        '(r["val_metrics"]["f1_false"], r["val_metrics"]["f1_macro"]))',
    )
    document.add_heading("Week 5: Sequential LSTM", level=2)
    document.add_paragraph(
        "The LSTM task uses the prior 20 calls within the same operating day "
        "to classify the next call. No sequence crosses a date boundary. A "
        "chronological future-period test is reserved; the scaler is fitted "
        "only on training-window features. The notebook compares six settings "
        "using validation balanced accuracy, followed by validation accuracy "
        "and loss. It uses a weighted binary loss to emphasize missed-standard "
        "examples. `wait_length` and fields known only after answering or "
        "finishing a call are excluded."
    )
    document.add_paragraph(
        f"The validation-selected setting used hidden size "
        f"{w5_selected['hidden_size']}, {w5_selected['num_layers']} LSTM layer, "
        f"learning rate {w5_selected['learning_rate']}, dropout "
        f"{w5_selected['dropout']}, and missed-class loss weight "
        f"{w5_selected['false_class_weight']}."
    )
    add_code(
        document,
        "best_experiment_name = max(experiment_results, key=lambda name: "
        "(experiment_results[name]['validation_balanced_accuracy'], "
        "experiment_results[name]['best_val_accuracy'], "
        "-experiment_results[name]['best_val_loss']))",
    )
    add_code(
        document,
        "self.lstm = nn.LSTM(input_size=input_size, hidden_size=hidden_size, "
        "num_layers=num_layers, batch_first=True)\n"
        "_, (hidden_state, _) = self.lstm(inputs)\n"
        "last_hidden_state = hidden_state[-1]\n"
        "logits = self.linear(self.dropout(last_hidden_state))\n"
        "return logits.squeeze(-1)",
    )
    document.add_heading("Week 6: Duplicate-Aware Multiclass Intent Models", level=2)
    document.add_paragraph(
        "The input is the customer's `instruction`; the target is `intent` "
        "among 27 classes. Instructions are case-folded, whitespace-normalized, "
        "and grouped before a seeded, stratified 70/15/15 split. This keeps "
        "normalized duplicate requests from appearing in both train and held-out "
        "sets. The TF-IDF vocabulary is learned from training data only. The "
        "two existing neural models are a feed-forward MLP over TF-IDF features "
        "and a PyTorch LSTM over token sequences. Macro precision, recall, F1, "
        "one-vs-rest macro ROC-AUC, confusion matrices, and per-intent results "
        "are saved."
    )
    document.add_paragraph(
        "The TF-IDF representation uses unigrams and bigrams with a maximum "
        "vocabulary of 12,000 terms. The MLP has 256- and 128-unit hidden "
        "layers. The separate token LSTM uses a 64-dimensional embedding, "
        "64-unit hidden state, and dropout 0.30."
    )
    add_code(
        document,
        'normalized_instructions = (df["instruction"].astype(str).str.casefold()'
        '.str.replace(r"\\s+", " ", regex=True).str.strip())',
    )
    add_code(
        document,
        "train_groups, holdout_groups = train_test_split(\n"
        "    all_group_ids, test_size=0.30, random_state=SEED, stratify=group_y\n"
        ")\n"
        "valid_groups, test_groups = train_test_split(\n"
        "    holdout_groups, test_size=0.50, random_state=SEED,\n"
        "    stratify=group_y[holdout_groups]\n"
        ")",
    )
    add_code(
        document,
        "X_train_tfidf = vectorizer.fit_transform(texts[train_idx])\n"
        "X_valid_tfidf = vectorizer.transform(texts[valid_idx])\n"
        "X_test_tfidf = vectorizer.transform(texts[test_idx])",
    )
    add_code(
        document,
        "self.network = nn.Sequential(\n"
        "    nn.Linear(input_size, 256), nn.ReLU(), nn.Dropout(0.35),\n"
        "    nn.Linear(256, 128), nn.ReLU(), nn.Dropout(0.25),\n"
        "    nn.Linear(128, number_of_classes),\n"
        ")\n"
        "return self.network(features)",
    )
    add_code(
        document,
        "self.embedding = nn.Embedding(vocab_size, embedding_size, padding_idx=0)\n"
        "self.lstm = nn.LSTM(embedding_size, hidden_size, batch_first=True)\n"
        "self.dropout = nn.Dropout(0.30)\n"
        "self.output = nn.Linear(hidden_size, number_of_classes)\n"
        "return self.output(self.dropout(hidden[-1]))",
    )
    document.add_paragraph(
        "The code excerpts above represent the existing model-selection and "
        "duplicate-normalization logic; full implementations and execution "
        "details remain in the linked notebooks and experiment runner."
    )

    document.add_heading("Held-Out Evaluation Results", level=1)
    document.add_paragraph(
        "Precision, recall, F1, and AUC are reported with macro averaging for "
        "class-comparable summaries. For the binary call-centre tasks, ROC-AUC "
        "is for the positive class `meets_standard=True`. Week 6 reports "
        "one-vs-rest macro ROC-AUC. The datasets, targets, and splits differ, "
        "so cross-stage scores should not be interpreted as a direct model "
        "competition."
    )
    add_table(
        document,
        ["Experiment", "Accuracy", "Macro precision", "Macro recall", "Macro F1", "ROC-AUC"],
        [
            [
                "Week 4 selected " + w4["best_experiment_name"],
                fmt(w4m["accuracy"]),
                fmt(w4m["precision_macro"]),
                fmt(w4m["recall_macro"]),
                fmt(w4m["f1_macro"]),
                fmt_auc(w4m["roc_auc_positive_meets_standard"]),
            ],
            [
                "Week 5 selected LSTM",
                fmt(w5m["accuracy"]),
                fmt(w5macro["precision"]),
                fmt(w5macro["recall"]),
                fmt(w5macro["f1-score"]),
                fmt_auc(w5m["roc_auc"]),
            ],
            [
                "Week 6 TF-IDF MLP",
                fmt(w6_mlp["accuracy"]),
                fmt(w6_mlp["precision_macro"]),
                fmt(w6_mlp["recall_macro"]),
                fmt(w6_mlp["f1_macro"]),
                fmt_auc(w6_mlp["roc_auc_ovr_macro"]),
            ],
            [
                "Week 6 PyTorch LSTM",
                fmt(w6_lstm["accuracy"]),
                fmt(w6_lstm["precision_macro"]),
                fmt(w6_lstm["recall_macro"]),
                fmt(w6_lstm["f1_macro"]),
                fmt_auc(w6_lstm["roc_auc_ovr_macro"]),
            ],
        ],
    )
    add_figure(
        document,
        ROOT / "week 4" / "images" / "fig8_confusion_matrix.png",
        "Week 4: selected model's chronological held-out confusion matrix.",
    )
    add_figure(
        document,
        ROOT / "week 4" / "images" / "fig10_roc_curve.png",
        "Week 4: ROC curve for the held-out positive class.",
    )
    add_figure(
        document,
        ROOT / "week 4" / "images" / "fig_train_exp2_capacity.png",
        "Week 4: selected capacity experiment's training and validation history.",
    )
    add_figure(
        document,
        ROOT / "week 5" / "images" / "fig9_test_results.png",
        "Week 5: future-period accuracy comparison and LSTM confusion matrix.",
    )
    add_figure(
        document,
        ROOT / "week 5" / "images" / "fig10_roc_curve.png",
        "Week 5: ROC curve for the future-period test set.",
    )
    add_figure(
        document,
        ROOT / "week 5" / "images" / "figure-01-training-loss-plots.png",
        "Week 5: training and validation loss histories for the compared LSTM configurations.",
    )
    add_figure(
        document,
        ROOT / "week 6" / "images" / "confusion_matrices.png",
        "Week 6: held-out confusion matrices for the MLP and LSTM.",
    )
    add_figure(
        document,
        ROOT / "week 6" / "images" / "roc_curves.png",
        "Week 6: one-vs-rest ROC curves across the 27 intent classes.",
    )

    document.add_heading("Performance Analysis and Limitations", level=1)
    document.add_paragraph(
        f"Week 4's selected model ({w4['best_experiment_name']}) achieved "
        f"{w4m['accuracy']:.1%} accuracy, macro F1 {w4m['f1_macro']:.3f}, "
        f"and positive-class ROC-AUC {w4m['roc_auc_positive_meets_standard']:.3f}. "
        f"For missed-standard calls, precision, recall, and F1 were "
        f"{w4m['precision_false']:.3f}, {w4m['recall_false']:.3f}, and "
        f"{w4m['f1_false']:.3f}; for calls meeting the standard, these were "
        f"{w4m['precision_true']:.3f}, {w4m['recall_true']:.3f}, and "
        f"{w4m['f1_true']:.3f}. The test "
        "period has a different missed-class prevalence from the training "
        "period, and its low accuracy and near-chance AUC indicate weak "
        "generalization; this should not be treated as an operational predictor."
    )
    document.add_paragraph(
        f"Week 5's selected LSTM achieved {w5m['accuracy']:.1%} test accuracy "
        f"versus {w5m['train_majority_baseline_accuracy']:.1%} for the "
        f"training-majority baseline and ROC-AUC {w5m['roc_auc']:.3f}. Its "
        "missed-standard precision, recall, and F1 were "
        f"{w5m['missed_standard']['precision']:.3f}, "
        f"{w5m['missed_standard']['recall']:.3f}, and "
        f"{w5m['missed_standard']['f1-score']:.3f}. The model therefore did "
        "not beat the baseline accuracy and captured few missed-standard calls. "
        "This is a substantive negative finding, not a success claim."
    )
    document.add_paragraph(
        f"Week 6's majority-class baseline accuracy was "
        f"{w6m['Majority baseline']['accuracy']:.3f}; the MLP achieved macro "
        f"F1 {w6_mlp['f1_macro']:.3f} and "
        f"{w6_lstm['f1_macro']:.3f} (LSTM), with high one-vs-rest AUC. "
        "Those results are based on a duplicate-aware split, but remain "
        "specific to a synthetic, automatically labeled dataset. Similar "
        "templated language, dataset construction, and domain coverage can "
        "make these scores substantially more favorable than real deployment "
        "performance."
    )
    document.add_paragraph(
        "Further limitations include simulation-specific call patterns, "
        "temporal distribution shift, the use of a single held-out period in "
        "the call-centre experiments, and the absence of external validation. "
        "A future operational study would require realistic prospective data, "
        "calibration and threshold analysis, and monitoring of changing queue "
        "conditions. The Week 2 stated wait-time regression objective remains "
        "a distinct future model-development task; Weeks 4 and 5 do not evaluate "
        "regression error."
    )

    document.add_heading("Reproducibility and Project Artifacts", level=1)
    document.add_paragraph(
        "The Week 4 experiment runner records seed 42 and writes its metrics "
        "and plots to `week 4/results.json` and `week 4/images/`. Week 5 records "
        "seed 42, selected settings, test metrics, confusion matrix, and "
        "classification report in `week 5/outputs/week5_results.json`; the "
        "notebook saves the evaluation figures in `week 5/images/`. Week 6 "
        "records its split metadata, metrics, confusion matrices, and per-intent "
        "results in `week 6/outputs/week6_results.json`. Run instructions and "
        "dependencies are maintained in the respective READMEs and requirements "
        "files."
    )
    document.add_paragraph(
        "To complete the previously missing Week 7 evaluation outputs, the "
        "existing Week 4 experiment runner and Week 5 notebook were executed "
        "without changing their dataset, target, split strategy, or model "
        "design; the runs generated the saved test AUC values, ROC plots, and "
        "machine-readable Week 5 results. Week 6 metrics are read from its "
        "existing duplicate-aware results artifact. The report presents the "
        "saved artifacts rather than estimating or fabricating scores."
    )
    document.add_paragraph(
        f"GitHub repository: {GITHUB_URL}\n"
        f"Call-centre dataset: {CALL_DATA_URL}\n"
        f"Bitext dataset: {BITEXT_DATA_URL}"
    )

    document.save(REPORT_PATH)
    print(f"Saved project-wide report: {REPORT_PATH}")


if __name__ == "__main__":
    main()
