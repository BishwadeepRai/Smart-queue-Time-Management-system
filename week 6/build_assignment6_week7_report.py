"""Build an editable Week 7 technical report from the Assignment 6 artifacts."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import nbformat
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WEEK_DIR = Path(__file__).resolve().parent
ROOT = WEEK_DIR.parent
RESULTS_PATH = WEEK_DIR / "outputs" / "week6_results.json"
NOTEBOOK_PATH = WEEK_DIR / "week6_customer_support_intent_classification.ipynb"
IMAGE_DIR = WEEK_DIR / "images"
ASSET_DIR = ROOT / "report_assets"
PIPELINE_PNG = ASSET_DIR / "assignment6_pipeline.png"
PIPELINE_SVG = ASSET_DIR / "assignment6_pipeline.svg"
REPORT_PATH = WEEK_DIR / "Week7_Assignment6_Technical_Report.docx"
DATASET_CARD_URL = "https://huggingface.co/datasets/bitext/Bitext-customer-support-llm-chatbot-training-dataset"
DATASET_FILE_URL = (
    DATASET_CARD_URL
    + "/resolve/main/Bitext_Sample_Customer_Support_Training_Dataset_27K_responses-v11.csv"
)
GITHUB_URL = "https://github.com/BishwadeepRai/Smart-queue-Time-Management-system"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, placeholder, end))


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, value) in enumerate(zip(header.cells, headers)):
        cell.text = str(value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[index])
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.line_spacing = 1
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8.5)
    for row_index, data in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, data)):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[index])
            if row_index % 2 == 1:
                set_cell_shading(cell, "F4F8FB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    document.add_paragraph()
    return table


def add_caption(document, number: int, title: str, note: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    label = paragraph.add_run(f"Figure {number}. ")
    label.bold = True
    paragraph.add_run(title).italic = True
    note_paragraph = document.add_paragraph()
    note_paragraph.paragraph_format.line_spacing = 1
    note_paragraph.paragraph_format.space_after = Pt(10)
    note_run = note_paragraph.add_run(f"Note. {note}")
    note_run.font.size = Pt(9)


def add_figure(document, path: Path, number: int, title: str, note: str, width=6.2) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Required report figure is missing: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(document, number, title, note)


def add_code(document, code: str, source: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    for line_number, line in enumerate(code.splitlines()):
        if line_number:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    source_line = document.add_paragraph()
    source_line.paragraph_format.line_spacing = 1
    source_line.paragraph_format.space_after = Pt(8)
    run = source_line.add_run(f"Source: {source}")
    run.italic = True
    run.font.size = Pt(9)


def extract_snippet(notebook, cell_index: int, start: str, end: str) -> str:
    source = "".join(notebook.cells[cell_index].source)
    first = source.index(start)
    last = source.index(end, first) if end else len(source)
    return source[first:last].rstrip()


def format_score(value, digits=4) -> str:
    return f"{value:.{digits}f}"


def format_percent(value) -> str:
    return f"{value * 100:.2f}%"


def format_pvalue_delta(value) -> str:
    return f"{value * 100:.3f} percentage points"


def make_pipeline_flowchart() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans"})
    fig, ax = plt.subplots(figsize=(13.5, 7.2))
    fig.patch.set_facecolor("#F8FAFC")
    ax.set_facecolor("#F8FAFC")
    ax.set_xlim(0, 13.5)
    ax.set_ylim(0, 7.2)
    ax.axis("off")
    ax.text(
        6.75, 6.82, "ASSIGNMENT 6 | CUSTOMER-SUPPORT INTENT CLASSIFICATION",
        ha="center", va="center", fontsize=15, fontweight="bold", color="#17324D"
    )

    color = "#5B2C83"
    light = "#F2EAF7"
    stages = [
        (0.3, 5.0, 2.25, 1.05, "1. Acquire & inspect", "Bitext CSV\n26,872 x 5 fields"),
        (2.85, 5.0, 2.25, 1.05, "2. EDA", "27 intents; quality,\nclass counts, length"),
        (5.4, 5.0, 2.25, 1.05, "3. Normalize & group", "Case/whitespace\nnormalized requests"),
        (7.95, 5.0, 2.25, 1.05, "4. Split groups", "Stratified 70/15/15;\nno group crossing"),
        (10.5, 5.0, 2.65, 1.05, "5. Fit representations", "TF-IDF + token vocab\nfit on train only"),
    ]
    for x, y, width, height, heading, body in stages:
        ax.add_patch(
            FancyBboxPatch(
                (x, y), width, height,
                boxstyle="round,pad=0.06,rounding_size=0.1",
                facecolor="white", edgecolor=color, linewidth=1.5,
            )
        )
        ax.text(x + width / 2, y + 0.72, heading, ha="center", va="center",
                fontsize=9, fontweight="bold", color=color)
        ax.text(x + width / 2, y + 0.32, body, ha="center", va="center",
                fontsize=8, color="#253746")

    for index in range(len(stages) - 1):
        x, y, width, _, _, _ = stages[index]
        nx = stages[index + 1][0]
        ax.add_patch(FancyArrowPatch((x + width + 0.02, y + 0.52), (nx - 0.02, y + 0.52),
                                     arrowstyle="-|>", mutation_scale=12, color=color, linewidth=1.4))

    branches = [
        (2.2, "TF-IDF feed-forward MLP", "12,000 max features; 256 / 128 units\nAdam 0.001; validation macro-F1 checkpoint"),
        (8.0, "Token-sequence LSTM", "15,000 max vocabulary; 48-token cap\n64 embedding / 64 hidden; validation macro-F1"),
    ]
    for x, heading, body in branches:
        ax.add_patch(
            FancyBboxPatch(
                (x, 2.8), 3.3, 1.15,
                boxstyle="round,pad=0.06,rounding_size=0.1",
                facecolor=light, edgecolor=color, linewidth=1.6,
            )
        )
        ax.text(x + 1.65, 3.55, heading, ha="center", va="center",
                fontsize=9.5, fontweight="bold", color=color)
        ax.text(x + 1.65, 3.12, body, ha="center", va="center", fontsize=8, color="#253746")

    for x in (5.2, 8.0):
        start_x = 11.83 if x < 6 else 10.8
        ax.add_patch(FancyArrowPatch((start_x, 5.0), (x + 1.65, 4.0),
                                     connectionstyle="arc3,rad=0.16" if x < 6 else "arc3,rad=-0.12",
                                     arrowstyle="-|>", mutation_scale=12, color=color, linewidth=1.4))

    ax.add_patch(
        FancyBboxPatch((4.45, 0.75), 4.6, 1.05,
                       boxstyle="round,pad=0.06,rounding_size=0.1",
                       facecolor="white", edgecolor="#1D6A86", linewidth=1.6)
    )
    ax.text(6.75, 1.48, "Held-out predictions & evaluation", ha="center", va="center",
            fontsize=10, fontweight="bold", color="#1D6A86")
    ax.text(6.75, 1.08, "Accuracy; macro P/R/F1; OvR macro AUC; per-intent metrics;\n"
            "row-normalized confusion matrices; micro-averaged ROC curves",
            ha="center", va="center", fontsize=8.5, color="#253746")
    for x in (3.85, 9.65):
        ax.add_patch(FancyArrowPatch((x, 2.78), (6.75, 1.83),
                                     connectionstyle="arc3,rad=0.10" if x < 6 else "arc3,rad=-0.10",
                                     arrowstyle="-|>", mutation_scale=12, color="#1D6A86", linewidth=1.4))
    ax.text(6.75, 0.3,
            "Selection uses validation data. The test split is held out for final evaluation.",
            ha="center", va="center", fontsize=9, color="#435466")

    fig.savefig(PIPELINE_PNG, dpi=220, bbox_inches="tight")
    fig.savefig(PIPELINE_SVG, bbox_inches="tight")
    plt.close(fig)


def configure_document(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.add_run("TECH 405 | CUSTOMER-SUPPORT INTENT CLASSIFICATION")
        header_run.font.name = "Times New Roman"
        header_run.font.size = Pt(8)
        header_run.font.color.rgb = RGBColor(90, 100, 110)
        add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(23, 50, 77)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)


def add_section_heading(document, number: int, title: str) -> None:
    document.add_heading(f"{number}. {title}", level=1)


def add_reference(document, text: str, url: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.45)
    paragraph.paragraph_format.first_line_indent = Inches(-0.45)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.add_run(text)
    if url:
        paragraph.add_run(" ")
        add_hyperlink(paragraph, url, url)


def build_report() -> None:
    if not RESULTS_PATH.is_file():
        raise FileNotFoundError(f"Saved results not found: {RESULTS_PATH}")
    if not NOTEBOOK_PATH.is_file():
        raise FileNotFoundError(f"Executed notebook not found: {NOTEBOOK_PATH}")
    for name in ("intent_class_distribution_and_duplicates.png", "training_loss.png",
                 "confusion_matrices.png", "roc_curves.png"):
        if not (IMAGE_DIR / name).is_file():
            raise FileNotFoundError(f"Saved experimental figure not found: {IMAGE_DIR / name}")

    results = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    notebook = nbformat.read(NOTEBOOK_PATH, as_version=4)
    make_pipeline_flowchart()

    metrics = results["metrics"]
    mlp = metrics["Feed-forward MLP"]
    lstm = metrics["PyTorch LSTM"]
    baseline = metrics["Majority baseline"]
    mlp_val = results["mlp_validation_macro_f1_by_epoch"]
    lstm_val = results["lstm_validation_macro_f1_by_epoch"]
    mlp_loss = results["mlp_training_loss"]
    lstm_loss = results["lstm_training_loss"]
    mlp_best_epoch = max(range(len(mlp_val)), key=mlp_val.__getitem__) + 1
    lstm_best_epoch = max(range(len(lstm_val)), key=lstm_val.__getitem__) + 1
    class_names = results["class_names"]
    cm_mlp = results["confusion_matrices"]["Feed-forward MLP"]
    cm_lstm = results["confusion_matrices"]["PyTorch LSTM"]
    mlp_correct = sum(cm_mlp[i][i] for i in range(len(class_names)))
    lstm_correct = sum(cm_lstm[i][i] for i in range(len(class_names)))
    mlp_errors = results["test_rows"] - mlp_correct
    lstm_errors = results["test_rows"] - lstm_correct
    delta_accuracy = mlp["accuracy"] - lstm["accuracy"]
    delta_f1 = mlp["f1_macro"] - lstm["f1_macro"]
    delta_auc = mlp["roc_auc_ovr_macro"] - lstm["roc_auc_ovr_macro"]
    per_intent = results["per_intent_metrics"]

    document = Document()
    configure_document(document)

    # 1. Title page
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    title.paragraph_format.line_spacing = 1.3
    run = title.add_run(
        "Customer-Support Intent Classification\n"
        "A Duplicate-Aware Evaluation of Neural Text Models"
    )
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    for line in (
        "Week 7 Technical Report",
        "Bishwadeep Rai",
        "Westcliff University",
        "TECH 405: Artificial Neural Network and Deep Learning",
        "Professor Aashish Dhakal",
        date(2026, 10, 18).strftime("%B %d, %Y"),
    ):
        p = document.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.add_run("Project: AI-Based Smart Queue & Time Management System")
    document.add_page_break()

    # 2. Introduction and problem statement
    add_section_heading(document, 2, "Introduction and Problem Statement")
    document.add_paragraph(
        "Customer-support requests arrive as free-form text and must be directed to a "
        "service function such as order cancellation, refund handling, delivery information, "
        "or account recovery. Manual triage can add handling time and route a request to an "
        "unsuitable queue. This Assignment 6 project examines whether a model can infer a "
        "support intent from the text of an incoming request and provide a candidate routing "
        "label."
    )
    document.add_paragraph(
        "Formally, this is supervised, single-label multiclass text classification. For an "
        "input instruction x, the model predicts one label y from a fixed set of 27 intent "
        "classes. The notebook uses only the `instruction` field as model input and `intent` "
        "as the target. The task is a prototype evaluation on a public hybrid synthetic "
        "dataset—not a deployed queue-routing system or a test of wait time, service-level "
        "performance, or customer outcomes."
    )
    document.add_paragraph(
        "The central methodological concern was repeated text: several source rows contain "
        "identical instructions after case and whitespace normalization. If copies of the "
        "same request appeared in both training and testing, metrics could reward memorized "
        "phrases rather than generalization. The implementation therefore groups normalized "
        "duplicates before the stratified split and checks group separation."
    )

    # 3. Objectives and task
    add_section_heading(document, 3, "Project Objectives and Classification Task")
    for item in (
        "Inspect the source file and characterize class balance, missingness, repeated text, and instruction length.",
        "Predict the labeled `intent` from `instruction` only, among the dataset's 27 classes.",
        "Compare an order-insensitive TF-IDF feed-forward MLP with a token-sequence PyTorch LSTM.",
        "Prevent exact case/whitespace duplicate requests from crossing train, validation, and test partitions.",
        "Fit text representations on training data only and reserve the grouped test partition for final evaluation.",
        "Report accuracy, macro precision, macro recall, macro F1, multiclass ROC-AUC, per-intent metrics, and confusion matrices.",
        "Interpret results in light of synthetic/automatic labeling and limits on external validity.",
    ):
        document.add_paragraph(item, style="List Bullet")
    document.add_paragraph(
        "The majority-class baseline predicts the most frequent training intent for every test "
        "request. With 27 near-balanced classes, it provides a useful low-information reference; "
        "it is not expected to be a strong routing model."
    )

    # 4. Dataset
    add_section_heading(document, 4, "Dataset Description and Source")
    document.add_paragraph(
        "The notebook loads the Bitext Customer Service Tagged Training Dataset from the "
        "Hugging Face dataset repository. The project artifact records 26,872 examples and "
        "five selected fields: `flags`, `instruction`, `category`, `intent`, and `response`. "
        "Only `instruction` is passed to the classifiers; `intent` is the prediction target. "
        "`category`, `response`, and `flags` are excluded from model features because they can "
        "directly reveal or shortcut the target and/or represent information not available as "
        "the intended input."
    )
    document.add_paragraph(
        "The dataset card describes the corpus as hybrid synthetic data generated for customer "
        "service/virtual-assistant use and declares the CDLA-Sharing-1.0 license. The notebook "
        "downloads the CSV to a user cache rather than committing the raw data. Dataset license "
        "and card terms should be reviewed before redistribution (Bitext, n.d.)."
    )
    add_table(
        document,
        ["Property", "Verified project value"],
        [
            ["Dataset", results["dataset"]],
            ["Rows / selected columns", f"{results['source_rows']:,} / 5"],
            ["Target", "`intent` (27 nominal classes)"],
            ["Model input", "`instruction` text only"],
            ["Intent supports", f"{results['class_count_min']}–{results['class_count_max']} source rows per class"],
            ["Dataset source", DATASET_CARD_URL],
            ["Raw-data storage", "Downloaded to a local cache; not stored in this repository"],
        ],
        widths=[1.8, 4.7],
    )
    document.add_paragraph("The source fields and their use in this project are:")
    add_table(
        document,
        ["Source field", "Role in the experiment"],
        [
            ["`flags`", "Audited as source metadata; not used as a model input."],
            ["`instruction`", "Sole text predictor provided to the classifiers."],
            ["`category`", "Audited metadata; excluded from model input to avoid a target shortcut."],
            ["`intent`", "Ground-truth target label; encoded into 27 class IDs."],
            ["`response`", "Associated response text; excluded because it is not available at routing time and could reveal the label."],
        ],
        widths=[1.55, 4.95],
    )
    document.add_paragraph("Target classes observed in the saved results:")
    class_rows = [
        [results["class_names"][i], results["class_names"][i + 9], results["class_names"][i + 18]]
        for i in range(9)
    ]
    add_table(
        document,
        ["Intent class", "Intent class", "Intent class"],
        class_rows,
        widths=[2.16, 2.16, 2.16],
    )

    # 5. EDA
    add_section_heading(document, 5, "Exploratory Data Analysis")
    document.add_paragraph(
        f"The audit found {results['source_rows']:,} rows, {results['number_of_intents']} intents, "
        f"and no missing or blank values in any of the five inspected source fields. The smallest "
        f"and largest intent classes contain {results['class_count_min']} and "
        f"{results['class_count_max']} examples, respectively; the target is therefore close to "
        "balanced rather than dominated by one class. That makes macro-averaged metrics "
        "appropriate for comparing performance across intents."
    )
    wc = results["instruction_word_count"]
    document.add_paragraph(
        f"Instruction length averages {wc['mean']:.2f} whitespace-separated words; median "
        f"{wc['median']:.0f}; interquartile range {wc['q25']:.0f}–{wc['q75']:.0f}; observed "
        f"range {wc['min']:.0f}–{wc['max']:.0f}. Median length varies by intent (the notebook "
        "records medians from 6 to 11 words). The EDA is descriptive and does not establish that "
        "request length causes or independently predicts an intent."
    )
    document.add_paragraph(
        f"Case- and whitespace-normalization produced {results['unique_instruction_groups']:,} "
        f"unique instruction groups. Among them, {results['repeated_instruction_groups']:,} "
        f"groups are repeated, adding {results['additional_repeated_instruction_rows']:,} "
        "duplicate rows beyond one representative per group. The notebook checks that "
        "normalized duplicates do not map to conflicting intent labels. This evidence motivated "
        "grouped rather than row-wise splitting."
    )
    add_figure(
        document,
        IMAGE_DIR / "intent_class_distribution_and_duplicates.png",
        1,
        "Intent support, missingness, repeated normalized requests, and instruction length",
        "This four-panel figure is the saved exploratory output from the executed notebook. "
        "Near-balanced label counts do not rule out linguistic or dataset-construction bias.",
    )

    # 6. Pipeline and flowchart
    add_section_heading(document, 6, "Machine-Learning Development Pipeline")
    document.add_paragraph(
        "Figure 2 represents the actual Assignment 6 process. The dataset is first inspected; "
        "normalized duplicates are identified and assigned as groups before splitting. Training "
        "text alone is used to learn TF-IDF and token vocabularies. The MLP and LSTM are trained "
        "as parallel alternatives, selected/checkpointed using validation macro F1, and evaluated "
        "on the same held-out test groups. Predictions are converted to 27-class outputs for "
        "metrics and error analysis."
    )
    add_figure(
        document,
        PIPELINE_PNG,
        2,
        "Assignment 6 intent-classification pipeline implemented in the notebook",
        "The diagram is a process illustration generated from verified notebook settings; "
        "it is not a model output. The displayed result charts that follow are saved experimental outputs.",
        width=6.55,
    )
    add_table(
        document,
        ["Pipeline stage", "Purpose in this implementation"],
        [
            ["Acquire / inspect", "Load the public CSV; verify schema, dimensions, labels, missingness, and source values."],
            ["EDA", "Describe intent counts, repeated normalized instructions, and request lengths."],
            ["Normalize / group", "Create case-folded, whitespace-collapsed keys; validate consistent intent per duplicate group."],
            ["Split", "Stratify unique groups into train/validation/test to prevent normalized copies crossing partitions."],
            ["Represent text", "Fit TF-IDF and the token vocabulary on training rows only; transform the other partitions."],
            ["Train / select", "Train two distinct neural architectures; retain validation-selected checkpoints."],
            ["Predict / evaluate", "Generate held-out class probabilities and labels; calculate metrics, confusion matrices, and ROC curves."],
            ["Interpret", "Compare results and document corpus, leakage, and external-validity limits."],
        ],
        widths=[1.5, 5.0],
    )

    # 7. Preprocessing and split
    add_section_heading(document, 7, "Data Preprocessing and Train/Test Splitting")
    document.add_paragraph(
        "Text is normalized for duplicate detection by Unicode-aware case-folding, collapsing "
        "runs of whitespace to a single space, and trimming. A factorized group ID represents "
        "each normalized instruction. Before a split is accepted, group summaries verify that "
        "each group has exactly one intent label; conflicting labels raise an explicit error. "
        "This normalization is used to establish groups, while the actual original instruction "
        "strings remain the model inputs."
    )
    document.add_paragraph(
        f"The split uses random seed {results['seed']}. Unique instruction groups are first split "
        "into 70% training and 30% held-out groups with stratification by intent. The held-out "
        "30% is split equally into validation and test, yielding approximate 70/15/15 proportions. "
        "Rows are mapped back to those groups; disjointness assertions check that no group ID is "
        "shared between partitions. Every intent is asserted present in all three partitions."
    )
    add_table(
        document,
        ["Partition", "Rows", "Normalized groups", "Approx. row share", "Per-intent support"],
        [
            ["Train", f"{results['train_rows']:,}", f"{results['train_groups']:,}",
             f"{results['train_rows']/results['source_rows']:.1%}", "664–718"],
            ["Validation", f"{results['validation_rows']:,}", f"{results['validation_groups']:,}",
             f"{results['validation_rows']/results['source_rows']:.1%}", "138–160"],
            ["Test", f"{results['test_rows']:,}", f"{results['test_groups']:,}",
             f"{results['test_rows']/results['source_rows']:.1%}",
             f"{results['test_class_count_min']}–{results['test_class_count_max']}"],
        ],
        widths=[1.15, 1.0, 1.3, 1.1, 1.95],
    )
    document.add_paragraph(
        "For the MLP, `TfidfVectorizer.fit_transform` is applied to training text and the fitted "
        "vectorizer only transforms validation and test text. For the LSTM, the token-frequency "
        "vocabulary is built by iterating over training indices only; unseen tokens map to the "
        "unknown token. This prevents the held-out vocabulary from influencing learned text "
        "representations. The model target encoder maps labels to indices; the split is stratified "
        "on intent labels."
    )
    document.add_paragraph(
        "Duplicate grouping addresses exact duplicates under the specified normalization. It does "
        "not demonstrate that paraphrases, shared templates, or other semantically similar "
        "synthetic requests are independent across partitions."
    )

    # 8. Architectures
    add_section_heading(document, 8, "Model Architectures and Implementation")
    document.add_heading("8.1 TF-IDF feed-forward MLP", level=2)
    document.add_paragraph(
        "The MLP is a fast lexical baseline that converts each instruction into a sparse TF-IDF "
        "vector of unigrams and bigrams. The vectorizer lowercases text, retains up to 12,000 "
        "features, removes terms with document frequency below two, and uses sublinear term "
        "frequency (scikit-learn developers, n.d.). The feed-forward network has dense layers "
        "of 256 and 128 units with ReLU "
        "activations and dropout 0.35 and 0.25; the output has 27 logits. Training uses Adam "
        "(learning rate 0.001, weight decay 0.0001) and cross-entropy. Sparse rows are converted "
        "to dense batches of 256 for neural-network updates. Training is capped at 10 epochs with "
        "patience 2, retaining the checkpoint with best validation macro F1."
    )
    document.add_heading("8.2 PyTorch token LSTM", level=2)
    document.add_paragraph(
        "The LSTM models token sequences and can represent word-order context. Tokenization uses "
        "a lowercase regular expression for letters, digits, and apostrophes. The vocabulary "
        "contains at most 15,000 training-derived tokens plus padding and unknown tokens. "
        "Sequences are truncated/padded to 48 tokens. An embedding layer maps token IDs into "
        "64-dimensional vectors; a one-layer LSTM has hidden size 64, followed by dropout 0.30 "
        "and a 27-logit linear output. Packed sequences use the true instruction lengths. "
        "Training uses Adam at 0.001, batch size 128, cross-entropy, up to 8 epochs, and patience "
        "2 on validation macro F1 (PyTorch contributors, n.d.)."
    )
    document.add_heading("8.3 Baseline", level=2)
    document.add_paragraph(
        "The baseline selects the majority label from training labels and predicts it for every "
        "test example. This provides a reference for the class-balanced task and helps check that "
        "the neural models improve beyond a constant prediction."
    )
    document.add_paragraph(
        f"Notebook runtime output records PyTorch {torch_version_from_notebook(notebook)}, CPU "
        f"execution, and seed {results['seed']}. No GPU timing, training-time comparison, or "
        "deployment latency benchmark is recorded in the saved artifacts."
    )

    # 9. Intermediate development steps
    add_section_heading(document, 9, "Intermediate Development Steps and Outputs")
    add_table(
        document,
        ["Step", "Recorded output"],
        [
            ["Source load", f"{results['source_rows']:,} rows x five source fields; {results['number_of_intents']} classes."],
            ["Quality checks", "Zero missing values and zero blank instructions/labels after trimming in the five audited columns."],
            ["Duplicate audit", f"{results['unique_instruction_groups']:,} normalized groups; {results['repeated_instruction_groups']:,} repeated groups; {results['additional_repeated_instruction_rows']:,} additional duplicate rows; labels agree within groups."],
            ["Partitioning", f"Train {results['train_rows']:,}/{results['train_groups']:,} rows/groups; validation {results['validation_rows']:,}/{results['validation_groups']:,}; test {results['test_rows']:,}/{results['test_groups']:,}."],
            ["MLP training history", f"{len(mlp_loss)} epochs; training loss {mlp_loss[0]:.4f} to {mlp_loss[-1]:.4f}; best validation macro F1 {max(mlp_val):.4f} at epoch {mlp_best_epoch}."],
            ["LSTM training history", f"{len(lstm_loss)} epochs; training loss {lstm_loss[0]:.4f} to {lstm_loss[-1]:.4f}; best validation macro F1 {max(lstm_val):.4f} at epoch {lstm_best_epoch}."],
            ["Held-out predictions", f"{results['test_rows']:,} requests; MLP correct {mlp_correct:,}, errors {mlp_errors}; LSTM correct {lstm_correct:,}, errors {lstm_errors}."],
        ],
        widths=[1.55, 4.95],
    )
    add_figure(
        document,
        IMAGE_DIR / "training_loss.png",
        3,
        "Training loss and validation macro F1 histories",
        "Saved Assignment 6 plot. The MLP and LSTM use different representations and have different epoch counts; "
        "the plotted histories show their own training/validation trajectories and should not be treated as "
        "a compute-time comparison.",
    )
    document.add_paragraph(
        "Both models reached high validation macro F1. On the held-out test set, each scored slightly "
        "below its own best validation macro F1. The observed gaps are small, but a single split on "
        "synthetic/template-like data cannot rule out overfitting to dataset-specific lexical patterns."
    )

    # 10. Code snippets
    add_section_heading(document, 10, "Critical Code Snippets and Explanations")
    document.add_paragraph(
        "The excerpts below are selected from the executed Assignment 6 notebook; surrounding "
        "training boilerplate is omitted. They are intended to explain the design decisions rather "
        "than reproduce the complete source."
    )
    duplicate_code = extract_snippet(notebook, 5, "normalized_instructions = (", "missing_counts =")
    document.add_heading("10.1 Normalize duplicates and validate group labels", level=2)
    add_code(document, duplicate_code, "week6_customer_support_intent_classification.ipynb, EDA/preparation cell")
    document.add_paragraph(
        "The same normalization defines duplicate groups. Aggregating each group and checking "
        "`label_count` prevents a normalized instruction with conflicting labels from being "
        "silently assigned to a split."
    )
    split_code = extract_snippet(notebook, 7, "all_group_ids = np.arange", "y_train, y_valid, y_test =")
    document.add_heading("10.2 Stratify unique groups, then map groups to rows", level=2)
    add_code(document, split_code, "week6_customer_support_intent_classification.ipynb, split cell")
    document.add_paragraph(
        "The code splits group IDs, not row indices, and then maps all rows from each group to a "
        "single partition. The assertions immediately following this excerpt verify partition "
        "disjointness and that every class appears in each split."
    )
    tfidf_code = extract_snippet(notebook, 9, "vectorizer = TfidfVectorizer(", "print(f\"Training-only TF-IDF vocabulary:")
    document.add_heading("10.3 Fit text representation on training data only", level=2)
    add_code(document, tfidf_code, "week6_customer_support_intent_classification.ipynb, MLP cell")
    document.add_paragraph(
        "`fit_transform` learns vocabulary and inverse-document-frequency statistics from training "
        "text; validation and test use `transform` so their content does not fit the representation."
    )
    mlp_code = extract_snippet(notebook, 9, "class SparseTextMLP", "mlp = SparseTextMLP")
    document.add_heading("10.4 Feed-forward neural architecture", level=2)
    add_code(document, mlp_code, "week6_customer_support_intent_classification.ipynb, MLP cell")
    document.add_paragraph(
        "The MLP operates on TF-IDF features through two ReLU/dropout hidden layers and returns "
        "one logit per intent."
    )
    lstm_code = extract_snippet(notebook, 11, "class IntentLSTM", "def make_loader")
    document.add_heading("10.5 Token-sequence LSTM architecture", level=2)
    add_code(document, lstm_code, "week6_customer_support_intent_classification.ipynb, LSTM cell")
    document.add_paragraph(
        "The embedding converts token IDs to vectors. Packed sequences avoid treating padding as "
        "meaningful input; the final hidden representation is passed through dropout and a "
        "27-class linear layer."
    )
    metrics_code = extract_snippet(notebook, 13, "def score_model", "results = pd.DataFrame")
    document.add_heading("10.6 Multiclass metrics", level=2)
    add_code(document, metrics_code, "week6_customer_support_intent_classification.ipynb, evaluation cell")
    document.add_paragraph(
        "Macro averaging gives each intent equal weight. The code computes multiclass one-vs-rest "
        "macro ROC-AUC from class probabilities; the ROC chart uses a separately flattened "
        "micro-averaged construction, as noted in the figure caption (scikit-learn developers, n.d.)."
    )

    # 11. Setup and metrics
    add_section_heading(document, 11, "Experimental Setup and Evaluation Metrics")
    document.add_paragraph(
        f"Experiments use random seed {results['seed']}. The recorded environment reports "
        f"PyTorch {torch_version_from_notebook(notebook)} on CPU. The notebook uses scikit-learn "
        "for TF-IDF, label encoding, split generation, classification metrics, and ROC calculations; "
        "PyTorch provides the neural architectures and training loops. Reproducibility is "
        "supported by a fixed seed and saved split/result artifacts, although exact training "
        "determinism can still depend on package versions and hardware."
    )
    add_table(
        document,
        ["Metric", "Interpretation in this task"],
        [
            ["Accuracy", "Share of test requests assigned the correct intent."],
            ["Macro precision", "Unweighted mean of per-intent precision; penalizes false routing per class equally."],
            ["Macro recall", "Unweighted mean of per-intent recall; shows class-balanced coverage."],
            ["Macro F1", "Unweighted mean of per-intent F1; primary validation checkpoint/selection measure."],
            ["ROC-AUC (OvR macro)", "Mean one-vs-rest ranking AUC across 27 classes using class probabilities; threshold-independent."],
            ["Confusion matrix", "Actual intent by predicted intent; saved plot is normalized by actual class (row-wise)."],
        ],
        widths=[1.7, 4.8],
    )
    document.add_paragraph(
        "Validation macro F1 was used to retain model checkpoints. The test partition was evaluated "
        "after model selection and reports the primary held-out comparison. ROC-AUC applies to "
        "probability scores; the hard-label majority baseline has no meaningful AUC here and is "
        "stored as not calculated."
    )

    # 12. Results and visual evidence
    add_section_heading(document, 12, "Results, Tables, and Visual Evidence")
    document.add_paragraph(
        "The table below contains the saved test-set metrics. Values are read directly from "
        "`outputs/week6_results.json`; no metrics were re-estimated for this report."
    )
    add_table(
        document,
        ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro F1", "OvR macro AUC"],
        [
            ["Majority baseline", *(format_score(baseline[k]) for k in
             ["accuracy", "precision_macro", "recall_macro", "f1_macro"]), "Not calculated"],
            ["Feed-forward MLP", *(format_score(mlp[k]) for k in
             ["accuracy", "precision_macro", "recall_macro", "f1_macro"]),
             format_score(mlp["roc_auc_ovr_macro"], 5)],
            ["PyTorch LSTM", *(format_score(lstm[k]) for k in
             ["accuracy", "precision_macro", "recall_macro", "f1_macro"]),
             format_score(lstm["roc_auc_ovr_macro"], 5)],
        ],
        widths=[1.45, 0.95, 1.0, 0.95, 0.8, 1.35],
    )
    document.add_paragraph("Held-out test-set performance (n = 4,031).")
    add_figure(
        document,
        IMAGE_DIR / "confusion_matrices.png",
        4,
        "Held-out row-normalized confusion matrices for the MLP and LSTM",
        "Rows are ground-truth intents and columns are predictions; each row sums to one. "
        "The MLP has 3,999 correct predictions and 32 errors; the LSTM has 3,984 correct "
        "predictions and 47 errors. The full integer matrices are saved in the JSON results.",
        width=6.45,
    )
    add_figure(
        document,
        IMAGE_DIR / "roc_curves.png",
        5,
        "Held-out micro-averaged one-vs-rest ROC curves",
        "The notebook flattens the binarized labels and class probabilities and plots micro-AUC. "
        "The table reports the separate macro one-vs-rest AUC values: 0.99996 for the MLP and "
        "0.99991 for the LSTM. Do not label this plotted micro curve as macro-averaged.",
        width=5.8,
    )
    document.add_paragraph(
        "The requested screenshots were not present as separate screenshot files in the provided "
        "project artifacts. Figures 1, 3, 4, and 5 are genuine saved charts generated by the "
        "executed notebook—not screenshots of a user interface. Capture any additional notebook "
        "or editor screenshots directly from the executed project before submission; do not "
        "present the pipeline diagram as experimental output."
    )
    add_table(
        document,
        ["Suggested screenshot filename", "Capture from the executed project", "Evidence to show"],
        [
            ["assignment6_dataset_eda.png", "Notebook EDA cells and saved EDA plot",
             "26,872 x 5 rows, 27 intents, missingness/duplicate audit, class counts."],
            ["assignment6_group_split.png", "Notebook group-splitting cell output",
             "Seeded train/validation/test row and group counts with disjointness assertions."],
            ["assignment6_model_training.png", "Notebook training histories and selected checkpoints",
             "MLP and LSTM train/validation macro-F1 history and selected epochs."],
            ["assignment6_final_test_results.png", "Notebook evaluation cell and result table",
             "Accuracy, macro P/R/F1, macro OvR AUC, per-intent results, confusion matrices, and ROC figure."],
        ],
        widths=[1.75, 2.2, 2.55],
    )
    document.add_paragraph(
        "These are capture recommendations only; they are not embedded or claimed as existing "
        "screenshots. Use screenshots of the actual executed notebook, keep the metric names and "
        "class labels visible, and do not edit numerical results."
    )

    # 13. Comparative performance
    add_section_heading(document, 13, "Comparative Performance Analysis")
    document.add_paragraph(
        f"The MLP exceeds the LSTM by {delta_accuracy * 100:.3f} percentage points in accuracy "
        f"and {delta_f1 * 100:.3f} percentage points in macro F1, with {mlp_correct-lstm_correct} "
        f"additional correct assignments on the same {results['test_rows']:,} test rows. Its "
        f"macro one-vs-rest AUC advantage is {delta_auc:.5f}. Both models substantially exceed "
        f"the {format_percent(baseline['accuracy'])} majority-baseline accuracy. The MLP is the "
        "stronger measured model on this test split, but the metric differences are small; the "
        "evidence does not establish a statistically significant or deployment-level advantage."
    )
    document.add_paragraph(
        "The observed outcome is consistent with the TF-IDF MLP being well matched to this corpus: "
        "unigrams and bigrams can encode distinctive intent phrases, while the task contains "
        "short, formulaic, automatically labeled requests. This is a plausible explanation, not "
        "a tested causal result. The LSTM's ability to model word order did not improve aggregate "
        "performance on this test partition; the MLP uses sparse TF-IDF and smaller batchwise "
        "dense conversion, while the LSTM uses train-derived token embeddings and sequence "
        "recurrent computation. Runtime, memory, and throughput were not measured, so computational "
        "cost cannot be ranked from the available evidence. This does not imply sequence models are generally "
        "inferior."
    )
    comparison_rows = []
    for intent in ("cancel_order", "change_order", "check_invoice", "get_invoice"):
        left = per_intent["Feed-forward MLP"][intent]
        right = per_intent["PyTorch LSTM"][intent]
        comparison_rows.append(
            [intent, str(int(left["support"])), format_score(left["precision"]),
             format_score(left["recall"]), format_score(left["f1-score"]),
             format_score(right["precision"]), format_score(right["recall"]),
             format_score(right["f1-score"])]
        )
    add_table(
        document,
        ["Intent", "n", "MLP P", "MLP R", "MLP F1", "LSTM P", "LSTM R", "LSTM F1"],
        comparison_rows,
        widths=[1.3, 0.42, 0.73, 0.73, 0.73, 0.75, 0.75, 0.75],
    )
    document.add_paragraph(
        "Per-intent values show that aggregate differences do not favor one architecture "
        "uniformly: the MLP has higher F1 for `change_order` and `check_invoice`, while the LSTM "
        "has higher F1 for `cancel_order` and `get_invoice`. These are observed class-level "
        "differences in this test set, not evidence of stable class-specific superiority."
    )
    document.add_paragraph(
        f"Validation-best macro F1 was {max(mlp_val):.4f} for the MLP and {max(lstm_val):.4f} "
        f"for the LSTM; corresponding test scores are {mlp['f1_macro']:.4f} and "
        f"{lstm['f1_macro']:.4f}. The validation-to-test gaps are "
        f"{max(mlp_val)-mlp['f1_macro']:.4f} and {max(lstm_val)-lstm['f1_macro']:.4f}. "
        "These small gaps are observed. They do not rule out overfitting to synthetic templates "
        "or answer whether performance transfers to naturally occurring user requests."
    )

    # 14. Limitations
    add_section_heading(document, 14, "Discussion of Limitations and Potential Improvements")
    document.add_heading("Observed limitations", level=2)
    for item in (
        "The dataset card describes hybrid synthetic/automatically labeled content. High scores measure agreement with those labels on this corpus, not correctness for real customer intents or actual queue-routing benefit.",
        "The split prevents exact normalized duplicates from crossing partitions, but paraphrase/template overlap and near-duplicate semantics were not comprehensively measured.",
        "There is no external validation set, independently labeled human sample, prospective customer traffic, or operational routing outcome in this experiment.",
        "Each intent has relatively similar support (950–1,000 source rows), but similar counts do not guarantee equal linguistic coverage, label quality, or fairness across customer groups and writing styles.",
        "No demographic or protected-group attributes are part of the audited model input/results; fairness across such groups is therefore unmeasured, not proven.",
        "Training loss falls substantially while validation macro F1 improves. Validation-test differences are modest, but available artifacts do not contain a controlled learning-curve analysis sufficient to dismiss overfitting.",
        "The notebook ran on CPU and records no elapsed time, peak memory, energy use, or inference latency. Computational cost cannot be quantitatively compared from the available outputs.",
        "The ROC figure is micro-averaged whereas the headline scalar AUC is macro one-vs-rest. Both are valid summaries but answer different averaging questions and must be labeled distinctly.",
    ):
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Potential improvements", level=2)
    for item in (
        "Evaluate on a separately collected, human-labeled set of real customer messages, with clear privacy and consent controls.",
        "Add template-aware and semantic near-duplicate detection or hold out templates/sources to test lexical generalization beyond exact normalization.",
        "Report uncertainty with repeated seeded group splits or group-aware cross-validation, while keeping a final external test untouched.",
        "Inspect per-intent errors and confusion pairs with domain reviewers; improve ambiguous label definitions and assess annotation agreement.",
        "Measure calibration, selective prediction/abstention, and the impact of routing errors; a high macro F1 alone does not define a safe routing threshold.",
        "Benchmark wall-clock training, inference throughput, and memory with recorded hardware/software settings before making efficiency claims.",
        "Compare the current neural models with a simple TF-IDF linear classifier and a suitable pretrained text encoder only as future experiments, preserving the existing results as the baseline.",
    ):
        document.add_paragraph(item, style="List Bullet")
    document.add_paragraph(
        "Potential causes above are labeled as explanations or recommendations. The report "
        "distinguishes them from the measured test metrics and notebook observations."
    )

    # 15. Conclusion
    add_section_heading(document, 15, "Conclusion")
    document.add_paragraph(
        "Assignment 6 implements a reproducible, duplicate-aware 27-class intent-classification "
        "experiment using a TF-IDF feed-forward MLP and a token-sequence LSTM. The 26,872-row "
        "Bitext corpus contains 27 near-balanced labels; the notebook normalizes and groups "
        "repeated requests, splits groups by intent, and learns text representations from "
        "training data only. On 4,031 held-out rows, the MLP achieved 0.9921 accuracy and 0.9921 "
        "macro F1, marginally exceeding the LSTM's 0.9883 accuracy and 0.9883 macro F1. Both "
        "substantially exceed the 0.0335 majority baseline on this dataset."
    )
    document.add_paragraph(
        "The results support the conclusion that both architectures fit the dataset's labeled "
        "intent patterns very well, with a small measured advantage for the MLP on this split. "
        "Because the source is hybrid synthetic and automatically labeled, exact duplicates are "
        "not the only possible source of similarity, and no real-customer or operational "
        "validation was performed, the results should be treated as a dataset-specific prototype "
        "rather than evidence of production readiness."
    )

    # 16. References
    add_section_heading(document, 16, "References")
    document.add_paragraph(
        "References below are limited to the dataset source, software documentation, and the "
        "project artifacts actually used. The dataset page does not identify a conventional "
        "journal publication for this CSV; no paper citation is inferred."
    )
    add_reference(
        document,
        "Bitext. (n.d.). Bitext—Customer service tagged training dataset for LLM-based virtual assistants "
        "[Data set]. Hugging Face. Retrieved October 18, 2026.",
        DATASET_CARD_URL,
    )
    add_reference(
        document,
        "Bishwadeep Rai. (2026). Customer support intent classification: Assignment 6 executed notebook, "
        "saved results, and figures [Source code and data analysis]. GitHub.",
        GITHUB_URL,
    )
    add_reference(
        document,
        "scikit-learn developers. (n.d.). TfidfVectorizer API reference. Retrieved October 18, 2026.",
        "https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html",
    )
    add_reference(
        document,
        "scikit-learn developers. (n.d.). roc_auc_score API reference. Retrieved October 18, 2026.",
        "https://scikit-learn.org/stable/modules/generated/sklearn.metrics.roc_auc_score.html",
    )
    add_reference(
        document,
        "PyTorch contributors. (n.d.). torch.nn.LSTM documentation. Retrieved October 18, 2026.",
        "https://pytorch.org/docs/stable/generated/torch.nn.LSTM.html",
    )
    add_reference(
        document,
        "PyTorch contributors. (n.d.). torch.nn.Embedding documentation. Retrieved October 18, 2026.",
        "https://pytorch.org/docs/stable/generated/torch.nn.Embedding.html",
    )

    # 17. Repository
    add_section_heading(document, 17, "GitHub Repository")
    paragraph = document.add_paragraph(
        "The Assignment 6 notebook, report, output metrics, and supporting figures are maintained in the project repository:"
    )
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph = document.add_paragraph()
    add_hyperlink(paragraph, GITHUB_URL, GITHUB_URL)
    document.add_paragraph(
        "Assignment 6 directory: week 6. The raw Bitext CSV is retrieved from its dataset source "
        "and cached locally; it is not represented as a tracked repository file."
    )

    document.save(REPORT_PATH)
    print(f"Saved editable technical report: {REPORT_PATH}")
    print(f"Saved Assignment 6 pipeline diagram: {PIPELINE_PNG}")


def torch_version_from_notebook(notebook) -> str:
    for cell in notebook.cells:
        for output in cell.get("outputs", []):
            text = "".join(output.get("text", [])) if output.get("output_type") == "stream" else ""
            if "PyTorch:" in text:
                for line in text.splitlines():
                    if "PyTorch:" in line:
                        return line.split("PyTorch:", 1)[1].split("|", 1)[0].strip()
    return "version not recorded"


if __name__ == "__main__":
    build_report()
